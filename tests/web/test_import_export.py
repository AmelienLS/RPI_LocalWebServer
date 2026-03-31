"""Tests for /export_serigraphie, /import_serigraphie and /import_serigraphie/confirm routes."""
import csv
import io
import json
import sqlite3

import openpyxl
import pytest

_COLUMNS = ["ref_ecran", "libelle", "pcb", "fab", "n_fab", "type", "n"]


# ---------------------------------------------------------------------------
# Export XLSX
# ---------------------------------------------------------------------------

def test_export_requires_login(client):
    response = client.get("/export_serigraphie")
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/")


def test_export_requires_admin(client, set_user_session):
    set_user_session(admin=False)
    response = client.get("/export_serigraphie")
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/index")


def test_export_returns_xlsx(client, set_user_session, add_serigraphie):
    set_user_session(admin=True)
    add_serigraphie(ref_ecran=700, libelle="Ecran export", fab="EX", type_="T1", n="001")

    response = client.get("/export_serigraphie")
    assert response.status_code == 200
    assert "spreadsheetml" in response.content_type

    wb = openpyxl.load_workbook(io.BytesIO(response.data))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    refs = [str(r[0]) for r in rows]
    assert "700" in refs


def test_export_xlsx_has_expected_columns(client, set_user_session):
    set_user_session(admin=True)
    response = client.get("/export_serigraphie")
    assert response.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(response.data))
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    assert headers == _COLUMNS


# ---------------------------------------------------------------------------
# Import CSV
# ---------------------------------------------------------------------------

def _make_csv(rows: list[dict]) -> bytes:
    """Build a CSV bytes payload from a list of dicts."""
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=["ref_ecran", "libelle", "pcb", "fab", "n_fab", "type", "n"])
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode("utf-8")


def _make_xlsx(rows: list[dict]) -> bytes:
    """Build an XLSX bytes payload from a list of dicts."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["ref_ecran", "libelle", "pcb", "fab", "n_fab", "type", "n"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_requires_login(client):
    response = client.post("/import_serigraphie", data={})
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/")


def test_import_requires_admin(client, set_user_session):
    set_user_session(admin=False)
    response = client.post("/import_serigraphie", data={})
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/index")


def test_import_csv_inserts_new_rows(client, set_user_session, test_db):
    set_user_session(admin=True)
    csv_data = _make_csv([
        {"ref_ecran": "800", "libelle": "Import CSV", "pcb": "", "fab": "IM", "n_fab": "", "type": "TC", "n": "010"},
    ])
    response = client.post(
        "/import_serigraphie",
        data={"import_file": (io.BytesIO(csv_data), "test.csv")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200

    with sqlite3.connect(test_db) as conn:
        row = conn.execute("SELECT libelle FROM serigraphie WHERE ref_ecran = 800").fetchone()
    assert row is not None
    assert row[0] == "Import CSV"


def test_import_xlsx_inserts_new_rows(client, set_user_session, test_db):
    set_user_session(admin=True)
    xlsx_data = _make_xlsx([
        {"ref_ecran": "801", "libelle": "Import XLSX", "pcb": "", "fab": "IX", "n_fab": "", "type": "TX", "n": "011"},
    ])
    response = client.post(
        "/import_serigraphie",
        data={"import_file": (io.BytesIO(xlsx_data), "test.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200

    with sqlite3.connect(test_db) as conn:
        row = conn.execute("SELECT libelle FROM serigraphie WHERE ref_ecran = 801").fetchone()
    assert row is not None
    assert row[0] == "Import XLSX"


def test_import_conflict_shows_comparison_page(client, set_user_session, add_serigraphie):
    set_user_session(admin=True)
    add_serigraphie(ref_ecran=802, libelle="Ecran existant", fab="AB", type_="T1", n="001")

    csv_data = _make_csv([
        {"ref_ecran": "802", "libelle": "Ecran modifié", "pcb": "", "fab": "AB", "n_fab": "", "type": "T1", "n": "001"},
    ])
    response = client.post(
        "/import_serigraphie",
        data={"import_file": (io.BytesIO(csv_data), "test.csv")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    body = response.data.decode("utf-8")
    assert "802" in body
    assert "Ecran existant" in body
    assert "Ecran modifié" in body


def test_import_invalid_format_returns_error(client, set_user_session):
    set_user_session(admin=True)
    response = client.post(
        "/import_serigraphie",
        data={"import_file": (io.BytesIO(b"data"), "test.txt")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    body = response.data.decode("utf-8")
    assert "non supporté" in body.lower() or "erreur" in body.lower() or "format" in body.lower()


# ---------------------------------------------------------------------------
# Confirm (conflict resolution)
# ---------------------------------------------------------------------------

def test_confirm_overwrites_selected(client, set_user_session, add_serigraphie, test_db):
    set_user_session(admin=True)
    add_serigraphie(ref_ecran=803, libelle="Ancien libellé", fab="AB", type_="T1", n="001")

    conflicts_json = json.dumps([
        {"ref_ecran": "803", "libelle": "Nouveau libellé", "pcb": "", "fab": "AB", "n_fab": "", "type": "T1", "n": "001"}
    ])
    response = client.post(
        "/import_serigraphie/confirm",
        data={"conflicts_json": conflicts_json, "overwrite": "803"},
    )
    assert response.status_code == 200

    with sqlite3.connect(test_db) as conn:
        row = conn.execute("SELECT libelle FROM serigraphie WHERE ref_ecran = 803").fetchone()
    assert row[0] == "Nouveau libellé"


def test_confirm_keeps_unselected(client, set_user_session, add_serigraphie, test_db):
    set_user_session(admin=True)
    add_serigraphie(ref_ecran=804, libelle="Libellé conservé", fab="AB", type_="T1", n="001")

    conflicts_json = json.dumps([
        {"ref_ecran": "804", "libelle": "Nouveau libellé", "pcb": "", "fab": "AB", "n_fab": "", "type": "T1", "n": "001"}
    ])
    # No "overwrite" field submitted -> keep existing
    response = client.post(
        "/import_serigraphie/confirm",
        data={"conflicts_json": conflicts_json},
    )
    assert response.status_code == 200

    with sqlite3.connect(test_db) as conn:
        row = conn.execute("SELECT libelle FROM serigraphie WHERE ref_ecran = 804").fetchone()
    assert row[0] == "Libellé conservé"


def test_confirm_overwrite_all(client, set_user_session, add_serigraphie, test_db):
    set_user_session(admin=True)
    add_serigraphie(ref_ecran=805, libelle="Ancien A", fab="AB", type_="T1", n="001")
    add_serigraphie(ref_ecran=806, libelle="Ancien B", fab="AB", type_="T1", n="002")

    conflicts_json = json.dumps([
        {"ref_ecran": "805", "libelle": "Nouveau A", "pcb": "", "fab": "AB", "n_fab": "", "type": "T1", "n": "001"},
        {"ref_ecran": "806", "libelle": "Nouveau B", "pcb": "", "fab": "AB", "n_fab": "", "type": "T1", "n": "002"},
    ])
    response = client.post(
        "/import_serigraphie/confirm",
        data={"conflicts_json": conflicts_json, "overwrite_all": "1"},
    )
    assert response.status_code == 200

    with sqlite3.connect(test_db) as conn:
        a = conn.execute("SELECT libelle FROM serigraphie WHERE ref_ecran = 805").fetchone()
        b = conn.execute("SELECT libelle FROM serigraphie WHERE ref_ecran = 806").fetchone()
    assert a[0] == "Nouveau A"
    assert b[0] == "Nouveau B"
